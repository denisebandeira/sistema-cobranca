"""Gera comunicados de cobrança a partir da Base Mestre.

O primeiro comunicado do arquivo modelo é usado como folha-base. O módulo
duplica essa folha para cada conjunto condomínio/unidade/condômino e preenche
as parcelas encontradas na aba ``Dados`` da Base Mestre.
"""

from __future__ import annotations

from copy import copy
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
import re
from typing import BinaryIO
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
from openpyxl import load_workbook


COLUNAS_OBRIGATORIAS = {
    "codigo_condominio",
    "condominio",
    "economia",
    "nome_condomino",
    "competencia",
    "data_vencimento",
    "vlr_corrigido",
}

COLUNAS_AGRUPAMENTO = [
    "codigo_condominio",
    "condominio",
    "economia",
    "nome_condomino",
]

LINHA_CABECALHO = 20
LINHA_INICIAL_PARCELAS = 21
LINHA_TOTAL_MODELO = 23
TAXA_HONORARIOS_PADRAO = 0.10


def carregar_base_mestre_excel(
    arquivo: str | Path | BinaryIO,
) -> pd.DataFrame:
    """Lê a aba Dados de uma Base Mestre gerada pelo sistema."""

    return pd.read_excel(
        arquivo,
        sheet_name="Dados",
        dtype={
            "codigo_condominio": "string",
            "economia": "string",
            "nosso_numero": "string",
        },
    )


def gerar_comunicados_excel(
    base_mestre: pd.DataFrame,
    modelo: str | Path | BinaryIO,
    *,
    data_referencia: date | datetime | None = None,
    taxa_honorarios: float = TAXA_HONORARIOS_PADRAO,
) -> bytes:
    """Retorna um XLSX com um comunicado por unidade/condômino.

    Somente parcelas vencidas há mais de 30 dias entram nos comunicados. Os
    honorários incidem sobre todas as parcelas exibidas, conforme o texto do
    modelo. A data de referência padrão é o dia da geração do arquivo.
    """

    base = _filtrar_base_elegivel(
        base_mestre,
        data_referencia=data_referencia,
        taxa_honorarios=taxa_honorarios,
    )
    return _gerar_xlsx_preparado(
        base,
        modelo,
        taxa_honorarios=taxa_honorarios,
    )


def gerar_comunicados_por_condominio(
    base_mestre: pd.DataFrame,
    modelo: str | Path | BinaryIO,
    *,
    data_referencia: date | datetime | None = None,
    taxa_honorarios: float = TAXA_HONORARIOS_PADRAO,
) -> dict[str, bytes]:
    """Retorna um XLSX separado para cada condomínio da Base Mestre."""

    base = _filtrar_base_elegivel(
        base_mestre,
        data_referencia=data_referencia,
        taxa_honorarios=taxa_honorarios,
    )
    arquivos: dict[str, bytes] = {}
    modelo_reutilizavel = _conteudo_modelo(modelo)

    grupos = base.groupby(
        ["codigo_condominio", "condominio"],
        sort=True,
        dropna=False,
    )
    for (codigo, condominio), dados_condominio in grupos:
        nome_arquivo = _nome_arquivo_condominio(str(codigo), str(condominio))
        arquivos[nome_arquivo] = _gerar_xlsx_preparado(
            dados_condominio,
            BytesIO(modelo_reutilizavel),
            taxa_honorarios=taxa_honorarios,
        )

    return arquivos


def gerar_pacote_comunicados_zip(
    base_mestre: pd.DataFrame,
    modelo: str | Path | BinaryIO,
    *,
    data_referencia: date | datetime | None = None,
    taxa_honorarios: float = TAXA_HONORARIOS_PADRAO,
) -> bytes:
    """Cria um ZIP contendo um arquivo Excel para cada condomínio."""

    arquivos = gerar_comunicados_por_condominio(
        base_mestre,
        modelo,
        data_referencia=data_referencia,
        taxa_honorarios=taxa_honorarios,
    )
    saida = BytesIO()
    with ZipFile(saida, mode="w", compression=ZIP_DEFLATED) as pacote:
        for nome, conteudo in arquivos.items():
            pacote.writestr(nome, conteudo)
    saida.seek(0)
    return saida.getvalue()


def _gerar_xlsx_preparado(
    base: pd.DataFrame,
    modelo: str | Path | BinaryIO,
    *,
    taxa_honorarios: float,
) -> bytes:
    workbook = load_workbook(modelo)
    folha_modelo = workbook.worksheets[0]

    for folha in list(workbook.worksheets[1:]):
        workbook.remove(folha)

    grupos = list(base.groupby(COLUNAS_AGRUPAMENTO, sort=True, dropna=False))
    folhas = [folha_modelo]
    for _ in grupos[1:]:
        folhas.append(workbook.copy_worksheet(folha_modelo))

    nomes_usados: set[str] = set()
    for folha, (chave, parcelas) in zip(folhas, grupos):
        codigo, condominio, economia, nome = (str(valor) for valor in chave)
        folha.title = _nome_aba_unico(
            codigo=codigo,
            economia=economia,
            nome=nome,
            usados=nomes_usados,
        )
        _preencher_comunicado(
            folha,
            parcelas=parcelas,
            condominio=condominio,
            economia=economia,
            nome=nome,
            taxa_honorarios=taxa_honorarios,
        )

    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except AttributeError:
        pass

    saida = BytesIO()
    workbook.save(saida)
    saida.seek(0)
    return saida.getvalue()


def _filtrar_base_elegivel(
    base_mestre: pd.DataFrame,
    *,
    data_referencia: date | datetime | None,
    taxa_honorarios: float,
) -> pd.DataFrame:
    if not 0 <= taxa_honorarios <= 1:
        raise ValueError("A taxa de honorários deve estar entre 0 e 1.")

    base = _preparar_base(base_mestre)
    referencia = _normalizar_data_referencia(data_referencia)
    limite = pd.Timestamp(referencia - timedelta(days=30))
    base = base.loc[base["data_vencimento"] < limite].copy()
    if base.empty:
        raise ValueError(
            "A Base Mestre não possui parcelas vencidas há mais de 30 dias."
        )
    return base


def _conteudo_modelo(modelo: str | Path | BinaryIO) -> bytes:
    if hasattr(modelo, "read"):
        modelo.seek(0)
        return modelo.read()
    return Path(modelo).read_bytes()


def _nome_arquivo_condominio(codigo: str, condominio: str) -> str:
    nome = re.sub(r"[\\/:*?\"<>|]", " ", condominio)
    nome = re.sub(r"\s+", " ", nome).strip(" .") or "condominio"
    return f"comunicados_{codigo}_{nome[:80]}.xlsx"


def gerar_comunicados_de_arquivo(
    base_mestre: str | Path | BinaryIO,
    modelo: str | Path | BinaryIO,
    *,
    data_referencia: date | datetime | None = None,
    taxa_honorarios: float = TAXA_HONORARIOS_PADRAO,
) -> bytes:
    """Atalho para ler a Base Mestre e gerar os comunicados."""

    dados = carregar_base_mestre_excel(base_mestre)
    return gerar_comunicados_excel(
        dados,
        modelo,
        data_referencia=data_referencia,
        taxa_honorarios=taxa_honorarios,
    )


def caminho_modelo_padrao() -> Path:
    """Localiza o modelo ao lado do código, inclusive no PyInstaller."""

    import sys

    if hasattr(sys, "_MEIPASS"):
        return Path(sys._MEIPASS) / "modelo_comunicado.xlsx"
    return Path(__file__).resolve().parent / "modelo_comunicado.xlsx"


def _preparar_base(base_mestre: pd.DataFrame) -> pd.DataFrame:
    ausentes = sorted(COLUNAS_OBRIGATORIAS - set(base_mestre.columns))
    if ausentes:
        raise ValueError(
            "A Base Mestre não contém as colunas obrigatórias: "
            + ", ".join(ausentes)
        )

    base = base_mestre.copy()
    for coluna in COLUNAS_AGRUPAMENTO + ["competencia"]:
        base[coluna] = base[coluna].fillna("").astype(str).str.strip()

    base["data_vencimento"] = pd.to_datetime(
        base["data_vencimento"],
        dayfirst=True,
        errors="coerce",
    )
    datas_invalidas = int(base["data_vencimento"].isna().sum())
    if datas_invalidas:
        raise ValueError(
            f"A Base Mestre possui {datas_invalidas} data(s) de vencimento inválida(s)."
        )

    base["vlr_corrigido"] = pd.to_numeric(
        base["vlr_corrigido"],
        errors="coerce",
    )
    valores_invalidos = int(base["vlr_corrigido"].isna().sum())
    if valores_invalidos:
        raise ValueError(
            f"A Base Mestre possui {valores_invalidos} valor(es) corrigido(s) inválido(s)."
        )

    grupos_vazios = base[COLUNAS_AGRUPAMENTO].eq("").any(axis=1)
    if grupos_vazios.any():
        raise ValueError(
            "A Base Mestre possui registros sem condomínio, unidade ou condômino."
        )

    return base.sort_values(
        COLUNAS_AGRUPAMENTO + ["data_vencimento", "competencia"],
        kind="stable",
    )


def _normalizar_data_referencia(valor: date | datetime | None) -> date:
    if valor is None:
        return date.today()
    if isinstance(valor, datetime):
        return valor.date()
    return valor


def _preencher_comunicado(
    folha,
    *,
    parcelas: pd.DataFrame,
    condominio: str,
    economia: str,
    nome: str,
    taxa_honorarios: float,
) -> None:
    quantidade = len(parcelas)
    linhas_extras = max(0, quantidade - 1)
    if linhas_extras:
        folha.insert_rows(LINHA_INICIAL_PARCELAS + 1, amount=linhas_extras)

    linha_final_parcelas = LINHA_INICIAL_PARCELAS + quantidade - 1
    linha_total = LINHA_TOTAL_MODELO + linhas_extras
    linha_honorarios = linha_total + 1
    linha_total_debito = linha_total + 2
    linha_observacao = linha_total + 3

    if taxa_honorarios == TAXA_HONORARIOS_PADRAO:
        descricao_taxa = "10% (dez por cento)"
    else:
        descricao_taxa = f"{taxa_honorarios:.2%}".replace(".00%", "%")

    texto = (
        "Olá!\n\n"
        "Referente:\n"
        f"Unidade {_unidade_exibicao(economia)}\n"
        f"Condomínio {_condominio_exibicao(condominio)}\n\n\n"
        "Meu nome é Carla e sou do escritório de cobrança.\n"
        "Gostaria de informar que a unidade acima encontra-se com pendências "
        "condominiais. Solicitamos a gentileza de entrar em contato conosco para "
        "regularização dessa pendência.\n"
        "Destacamos que, sobre as parcelas vencidas há mais de 30 dias, incidem "
        f"honorários de cobrança correspondentes a {descricao_taxa} do valor devido.\n"
        "*Favor desconsiderar esta comunicação caso o valor já tenha sido quitado.*\n"
        "Segue o demonstrativo do débito\n"
    )
    folha["B3"] = texto

    for deslocamento, (_, parcela) in enumerate(parcelas.iterrows()):
        linha = LINHA_INICIAL_PARCELAS + deslocamento
        if linha != LINHA_INICIAL_PARCELAS:
            _copiar_estilo_linha(
                folha,
                origem=LINHA_INICIAL_PARCELAS,
                destino=linha,
                coluna_inicial=3,
                coluna_final=5,
            )
        folha.row_dimensions[linha].height = 18
        folha.cell(linha, 3).value = parcela["competencia"]
        folha.cell(linha, 3).number_format = "@"
        folha.cell(linha, 4).value = parcela["data_vencimento"].to_pydatetime()
        folha.cell(linha, 4).number_format = "dd/mm/yyyy"
        folha.cell(linha, 5).value = float(parcela["vlr_corrigido"])
        folha.cell(linha, 5).number_format = '#,##0.00'

    folha.cell(linha_total, 3).value = "TOTAL PARCIAL"
    folha.cell(linha_total, 5).value = (
        f"=SUM(E{LINHA_INICIAL_PARCELAS}:E{linha_final_parcelas})"
    )

    taxa_excel = f"{taxa_honorarios:.10f}".rstrip("0").rstrip(".")
    folha.cell(linha_honorarios, 3).value = (
        f"HONORÁRIOS DE COBRANÇA ({taxa_honorarios:.0%})"
    )
    folha.cell(linha_honorarios, 5).value = f"=E{linha_total}*{taxa_excel}"
    folha.cell(linha_total_debito, 3).value = "TOTAL DO DÉBITO"
    folha.cell(linha_total_debito, 5).value = (
        f"=SUM(E{linha_total}:E{linha_honorarios})"
    )

    for linha in (linha_total, linha_honorarios, linha_total_debito):
        folha.row_dimensions[linha].height = 18
        folha.cell(linha, 5).number_format = '#,##0.00'

    folha.row_dimensions[linha_final_parcelas + 1].height = 18

    folha.cell(linha_observacao, 3).value = (
        "Taxa do Boleto R$ 1,99 a ser acrescida no valor total conforme o "
        "número de parcelas negociadas"
    )
    folha.row_dimensions[linha_observacao].height = 50.25
    folha.sheet_view.showGridLines = False
    folha.page_setup.orientation = "portrait"
    folha.page_setup.fitToWidth = 1
    folha.page_setup.fitToHeight = 0
    folha.sheet_properties.pageSetUpPr.fitToPage = True
    folha.print_area = f"A1:G{linha_observacao}"
    folha.print_title_rows = f"{LINHA_CABECALHO}:{LINHA_CABECALHO}"
    folha.print_options.horizontalCentered = True


def _copiar_estilo_linha(
    folha,
    *,
    origem: int,
    destino: int,
    coluna_inicial: int,
    coluna_final: int,
) -> None:
    for coluna in range(coluna_inicial, coluna_final + 1):
        celula_origem = folha.cell(origem, coluna)
        celula_destino = folha.cell(destino, coluna)
        celula_destino._style = copy(celula_origem._style)
        celula_destino.number_format = celula_origem.number_format
        celula_destino.protection = copy(celula_origem.protection)


def _unidade_exibicao(economia: str) -> str:
    unidade = re.sub(r"\s*\(\d+\)\s*$", "", economia).strip()
    correspondencia = re.fullmatch(r"([^\-]+)\s+-\s+\d+", unidade)
    if correspondencia:
        unidade = correspondencia.group(1).strip()
    return unidade


def _condominio_exibicao(condominio: str) -> str:
    if condominio.upper() == "CONJ. RES. HUMAITA GARDEN PARK":
        return "Garden Park Humaitá"
    return condominio


def _nome_aba_unico(
    *,
    codigo: str,
    economia: str,
    nome: str,
    usados: set[str],
) -> str:
    unidade = _unidade_exibicao(economia)
    base = f"{nome} {unidade}".strip()
    base = re.sub(r"[\\/*?:\[\]]", " ", base)
    base = re.sub(r"\s+", " ", base).strip(" '") or f"Comunicado {codigo}"
    base = base[:31].rstrip()

    candidato = base
    contador = 2
    while candidato.casefold() in usados:
        sufixo = f" ({contador})"
        candidato = f"{base[:31 - len(sufixo)]}{sufixo}"
        contador += 1

    usados.add(candidato.casefold())
    return candidato
